"""
모니터링으로 순위요약 엑셀이 갱신될 때마다 '모니터링완료_YYYY-MM-DD_HHMM.xlsx' 사본을 만듭니다. (디렉터 님 지시 2026-10-06)
  - 위치: 01.뷰성형외과 폴더 바로 아래  /  HHMM = 엑셀에 기록된 오늘 마지막 모니터링 시각(정규·수동 회차 또는 재확인)
  - 최신 1개만 유지 — 새 사본을 만들면 이전 날짜·시각 사본은 지움 (2026-10-12 디렉터 님 지시: 폴더 깊이 들어가지 않고 바로 확인)
  - 원본 순위요약_YYYY-MM.xlsx 는 그대로 두고 복사만 함
  python publish_done.py
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OUT = HERE.parent.parent if HERE.parent.parent.name == "01.뷰성형외과" else HERE / "모니터링완료"   # 01.뷰성형외과\


def publish(day=None):
    day = day or f"{datetime.now():%Y-%m-%d}"
    from paths import XL_DIR
    src = XL_DIR / f"순위요약_{day[:7]}.xlsx"
    if not src.exists():
        return f"{src.name} 없음"
    wb = load_workbook(src, read_only=True)
    last = max((str(r[0]) for ws in wb for r in ws.iter_rows(min_row=2, values_only=True)
                if r and str(r[0]).startswith(day)), default=None)
    wb.close()
    if not last:
        return f"{day} 모니터링 기록 없음"
    OUT.mkdir(exist_ok=True)
    name = f"모니터링완료_{day}_{last[11:13]}{last[14:16]}.xlsx"
    dst = OUT / name
    shutil.copy2(src, dst)
    for old in OUT.glob("모니터링완료_*.xlsx"):
        if old.name != name:
            try:
                old.unlink()
            except OSError:
                pass                                    # 열려 있으면 남겨 둠
    return f"{OUT.name}\\{name} 저장"


if __name__ == "__main__":
    print(publish(sys.argv[1] if len(sys.argv) > 1 else None))
