"""
오늘 최신 순위표(파워링크·파워컨텐츠 × PC·MO)를 순위요약_YYYY-MM.xlsx 에서 만들어 마크다운 표로 돌려줍니다.
  - 기준: 오늘 마지막 정규/수동 회차(재확인 행 제외) + 그 뒤 재확인에서 노출로 바뀐 값(★) 덮어쓰기
  - 에이전트 리더 대시보드는 '|' 로 시작하는 줄을 표로 그립니다 (디렉터 님 지시 2026-10-06)
  python rank_table.py           # 표 출력
  python rank_table.py --post    # 대시보드(뷰)순위담당 대화)에 표 게시
"""
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
from paths import XL_DIR, BOARD
SHEETS = ("파워링크", "파워컨텐츠")


def _is_recheck(r):
    return "재확인" in f"{r[4] or ''}{r[5] or ''}"


def latest(day=None):
    """→ (기준 회차 시각, 마지막 재확인 시각, {키워드: {(시트, 'PC'|'MO'): 표시값}}, 키워드 순서)"""
    day = day or f"{datetime.now():%Y-%m-%d}"
    path = XL_DIR / f"순위요약_{day[:7]}.xlsx"
    wb = load_workbook(path, read_only=True)
    data, order, base_ts, last_re = {}, [], None, None
    if "순위" in wb.sheetnames:                     # 2026-10-13: 한 시트(실행시각|구분|키워드|PC|MO|PC비고|MO비고)
        allrows = [r for r in wb["순위"].iter_rows(min_row=2, values_only=True) if r and str(r[0]).startswith(day)]
        by = {sh: [(r[0], *r[2:7]) for r in allrows if r[1] == sh] for sh in SHEETS}
    else:                                            # 예전 2시트 형식
        by = {sh: [r for r in wb[sh].iter_rows(min_row=2, values_only=True) if r and str(r[0]).startswith(day)]
              for sh in SHEETS if sh in wb.sheetnames}
    for sh, rows in by.items():
        runs = sorted({r[0] for r in rows if not _is_recheck(r)})
        if not runs:
            continue
        ts = runs[-1]
        base_ts = max(base_ts or ts, ts)
        for r in rows:
            if r[0] == ts and not _is_recheck(r):
                if r[1] not in order:
                    order.append(r[1])
                d = data.setdefault(r[1], {})
                d[(sh, "PC")], d[(sh, "MO")] = _fmt(r[2], r[4]), _fmt(r[3], r[5])
        for r in rows:
            if r[0] > ts and _is_recheck(r) and r[1] in data:
                last_re = max(last_re or r[0], r[0])
                if r[2] is not None:
                    data[r[1]][(sh, "PC")] = f"{r[2]}★"
                if r[3] is not None:
                    data[r[1]][(sh, "MO")] = f"{r[3]}★"
    wb.close()
    return base_ts, last_re, data, order


def _fmt(v, note):
    if v is None or v == "":
        return "-"
    note = str(note or "")
    return f"{v}(더보기)" if "더보기" in note else str(v)


def markdown(day=None):
    base_ts, last_re, data, order = latest(day)
    if not data:
        return "오늘 순위 결과 없음"
    head = f"■ 최신 순위표 (기준 {base_ts[5:16]} 회차" + (f", 재확인 {last_re[11:16]}까지 반영" if last_re else "") + ")"
    lines = [head, "| 키워드 | 파워링크 PC | 파워링크 MO | 파워컨텐츠 PC | 파워컨텐츠 MO |", "|---|---|---|---|---|"]
    for kw in order:
        d = data[kw]
        lines.append(f"| {kw} | " + " | ".join(d.get((sh, dv), "-") for sh in SHEETS for dv in ("PC", "MO")) + " |")
    lines.append("※ ★ = 재확인에서 새로 노출 확인 / (더보기) = 첫 화면 밖 더보기 목록 순위 / - = 해당 광고 대상 아님")
    return "\n".join(lines)


def post(day=None, status="done"):
    text = markdown(day)
    env = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
    subprocess.run([sys.executable, str(BOARD), "msg", "rank", text, "--status", status],
                   capture_output=True, text=True, env=env, timeout=60)
    return text


if __name__ == "__main__":
    print(post() if "--post" in sys.argv else markdown())
