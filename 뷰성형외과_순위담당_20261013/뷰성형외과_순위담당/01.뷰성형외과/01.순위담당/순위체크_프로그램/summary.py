"""
실행 결과를 '순위요약_YYYY-MM.xlsx' 에 세로로 누적합니다. (rank.run 이 CSV 저장 직후 호출)
  시트: '순위' 1개 — 회차마다 파워링크 키워드(settings.json keywords 순) 바로 아래 파워컨텐츠 키워드(pc_keywords 순)
  열:   실행시각 | 구분 | 키워드 | PC순위 | MO순위 | PC비고 | MO비고
  (2026-10-13 디렉터 님 지시: 시트 분리 없이 한 시트에서 PC·MO 열을 그대로 복사·붙여넣기)
순위 칸은 숫자 또는 '미노출'/'영역없음'. 색상 표시 없음 (디렉터 님 지시 2026-10-02, 첫 화면/더보기 구분은 비고 열).
"""
import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SHEETS = OrderedDict([("powerlink", "파워링크"), ("powercontents", "파워컨텐츠")])
SHEET = "순위"
HEADER = ["실행시각", "구분", "키워드", "PC순위", "MO순위", "PC비고", "MO비고"]
WIDTHS = [18, 11, 16, 9, 9, 22, 22]
TYPE_OF = {v: k for k, v in SHEETS.items()}                # '파워링크' → 'powerlink'
CENTER = Alignment(horizontal="center", vertical="center")


def summarize(rows):
    """CSV용 rows → {ad_type: {keyword: {device: (value, kind, note)}}}"""
    out = {t: OrderedDict() for t in SHEETS}
    for r in rows:
        t, kw, dev = r["ad_type"], r["keyword"], r["device"]
        if t not in out:
            continue
        cell = out[t].setdefault(kw, {})
        cur = cell.get(dev)
        if r["source"] == "1page":
            if r["total"] == -1:                       # 파워컨텐츠 영역 자체가 없음
                cell[dev] = ("영역없음", "none", "")
            elif r["target_hit"] == "Y":
                cell[dev] = (int(r["rank"]), "1page", f"첫 화면 {r['total']}개 중")
            elif cur is None:
                cell[dev] = ("미노출", "none", f"첫 화면 {r['total']}개")
        elif r["source"] == "more" and (cur is None or cur[1] == "none"):
            if r["target_hit"] == "Y":
                total = f"(전체 {r['total']}개)" if r["total"] else ""
                cell[dev] = (int(r["rank"]), "more", f"더보기 목록 {total}".strip())
            else:
                cell[dev] = ("미노출", "none", f"더보기 {r['total']}개까지 확인")
    return out


def keyword_order(base_dir: Path):
    """settings.json 의 키워드 순서 → {ad_type: [keyword, ...]}. 시트·붙여넣기 출력이 이 순서를 따른다."""
    try:
        s = json.loads((base_dir / "settings.json").read_text(encoding="utf-8"))
    except Exception:
        return {t: [] for t in SHEETS}
    return {"powerlink": [k.strip() for k in s.get("keywords", []) if k.strip()],
            "powercontents": [k.strip() for k in s.get("pc_keywords", []) if k.strip()]}


def reorder(data, base_dir: Path, fill=False):
    """summarize() 결과의 키워드를 settings.json 순서로 정렬. 설정에 없는 키워드는 뒤에 수집 순서대로.
    fill=True 면 설정에 있는데 결과가 없는 키워드도 빈 칸으로 넣어 붙여넣기 줄 수를 항상 맞춘다."""
    order = keyword_order(base_dir)
    for t in data:
        if fill and data[t]:
            for k in order.get(t, []):
                data[t].setdefault(k, {})
        pref = [k for k in order.get(t, []) if k in data[t]]
        rest = [k for k in data[t] if k not in pref]
        data[t] = OrderedDict((k, data[t][k]) for k in pref + rest)
    return data


def _sheet(wb, name=SHEET):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    ws.append(HEADER)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = CENTER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    return ws


def show_all(wb):
    """저장 전 정리 (2026-10-06): 남아 있는 필터 조건·숨김 행을 풀고 필터 범위를 전체로, 첫 행 고정.
    필터가 걸린 채 저장돼 최신 행이 안 보이던 문제 대응. 저장 후 첫 화면은 최신 행이 보이도록 맨 아래로 이동"""
    for ws in wb.worksheets:
        ws.auto_filter.filterColumn.clear()
        ws.auto_filter.sortState = None
        for d in ws.row_dimensions.values():
            d.hidden = False
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{max(ws.max_row, 1)}"
        ws.freeze_panes = "A2"
        last = f"A{max(ws.max_row, 2)}"
        ws.sheet_view.selection[0].activeCell = last
        ws.sheet_view.selection[0].sqref = last
        ws.sheet_view.topLeftCell = f"A{max(ws.max_row - 30, 2)}"


def _append_styled(ws, row):
    ws.append(list(row))
    r = ws.max_row
    ws.cell(r, 1).alignment = CENTER
    for col in (2, 4, 5):
        ws.cell(r, col).alignment = CENTER


def rows_for(data, ts):
    """summarize/reorder 결과 → 엑셀 행들. 파워링크 전체 다음 파워컨텐츠 전체 (한 회차가 이어진 블록)"""
    out = []
    for t, name in SHEETS.items():
        for kw, devs in data.get(t, {}).items():
            pc, mo = devs.get("pc"), devs.get("mobile")
            out.append([ts[:16], name, kw, pc[0] if pc else "", mo[0] if mo else "",
                        pc[2] if pc else "", mo[2] if mo else ""])
    return out


def old_rows(wb):
    """예전 형식(파워링크/파워컨텐츠 시트 2개) 행 → 새 형식 행. 회차(실행시각)별로 파워링크 → 파워컨텐츠 순"""
    rows = []
    names = list(SHEETS.values())
    for name in names:
        if name in wb.sheetnames:
            for i, r in enumerate(wb[name].iter_rows(min_row=2, values_only=True)):
                if r and r[0] is not None:
                    rows.append((str(r[0]), names.index(name), i, [r[0], name, *r[1:6]]))
    rows.sort(key=lambda x: x[:3])
    return [x[3] for x in rows]


def migrate(wb):
    """예전 2시트 → '순위' 1시트로 변환 (이미 변환됐으면 그대로). 변환하면 True"""
    if not any(n in wb.sheetnames for n in SHEETS.values()):
        return False
    ws = _sheet(wb)
    for row in old_rows(wb):
        _append_styled(ws, row)
    for name in SHEETS.values():
        if name in wb.sheetnames:
            wb.remove(wb[name])
    return True


def src_rows(src):
    """다른 통합문서(엑셀이 열려 있어 따로 저장된 파일 등)의 행을 새 형식으로"""
    if SHEET in src.sheetnames:
        return [list(r) for r in src[SHEET].iter_rows(min_row=2, values_only=True) if r and r[0] is not None]
    return old_rows(src)


def append_summary(rows, ts: str, base_dir: Path, log=print):
    """rows 를 요약해 순위요약_YYYY-MM.xlsx 에 추가. 파일이 엑셀에서 열려 있으면 임시 이름으로 저장."""
    if not rows:
        return None
    from paths import XL_DIR                    # 2026-10-06: 엑셀은 상위(01.뷰성형외과/01.순위담당) 폴더
    path = XL_DIR / f"순위요약_{ts[:7]}.xlsx"
    data = reorder(summarize(rows), base_dir, fill=True)
    new_rows = rows_for(data, ts)
    if path.exists():
        wb = load_workbook(path)
        migrate(wb)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    # 이전에 엑셀이 열려 있어 따로 저장된 파일(순위요약_YYYY-MM_HHMM.xlsx)이 있으면 먼저 합친다
    leftovers = sorted(XL_DIR.glob(f"{path.stem}_*.xlsx"))
    for lf in leftovers:
        try:
            src = load_workbook(lf)
            ws = _sheet(wb)
            for row in src_rows(src):
                _append_styled(ws, row)
            src.close()
            lf.unlink()
            log(f"{lf.name} 을 {path.name} 에 합쳤습니다.")
        except Exception as e:
            log(f"[!] {lf.name} 합치기 실패: {e}")
    ws = _sheet(wb)
    for row in new_rows:
        _append_styled(ws, row)
    try:
        show_all(wb)
        wb.save(path)
    except PermissionError:
        # 엑셀에서 열어둔 상태 → 이번 실행분만 따로 저장. 다음 실행 때 자동으로 합쳐진다.
        alt = path.with_name(path.stem + f"_{ts[11:16].replace(':', '')}" + path.suffix)
        one = Workbook()
        one.remove(one.active)
        ws = _sheet(one)
        for row in new_rows:
            _append_styled(ws, row)
        one.save(alt)
        log(f"[!] {path.name} 이 엑셀에서 열려 있어 이번 결과를 {alt.name} 으로 저장했습니다. "
            f"엑셀을 닫아두면 다음 실행 때 자동으로 {path.name} 에 합쳐집니다.")
        return alt
    log(f"엑셀 요약 저장: {path}")
    return path


if __name__ == "__main__":
    # python summary.py → 당월 순위요약 엑셀을 '순위' 1시트 형식으로 변환 (예전 2시트 파일용)
    from datetime import datetime
    from paths import XL_DIR
    path = XL_DIR / f"순위요약_{datetime.now():%Y-%m}.xlsx"
    wb = load_workbook(path)
    if migrate(wb):
        show_all(wb)
        wb.save(path)
        print(f"{path.name} → '{SHEET}' 시트 1개로 변환 ({wb[SHEET].max_row - 1}행)")
    else:
        print(f"{path.name} 은 이미 '{SHEET}' 시트 형식")
