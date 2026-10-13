"""
재확인(recheck_YYYYMMDD.csv)에서 '노출'로 바뀐 조합을 순위요약_YYYY-MM.xlsx 에 새 행으로 추가합니다.
  python recheck_excel.py              # 오늘
  python recheck_excel.py 20261006     # 특정 날짜
- 기존 행은 건드리지 않고, 재확인 시각(checked_at)으로 행 추가. 비고에 '재확인' 표시
- 이미 넣은 (재확인 시각, 키워드, 매체) 는 다시 넣지 않음 → 여러 번 실행해도 중복 없음
- 엑셀이 열려 있으면 summary.py 와 같이 순위요약_YYYY-MM_HHMM.xlsx 로 따로 저장 (다음 정규 실행 때 자동 합침)
디렉터 님 지시(2026-10-06): 뷰성형외과 순위 모니터링 결과를 원래 엑셀에 반영
"""
import csv
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from summary import SHEET, SHEETS, TYPE_OF, _append_styled, _sheet, migrate

HERE = Path(__file__).resolve().parent


def main():
    day = sys.argv[1] if len(sys.argv) > 1 else f"{datetime.now():%Y%m%d}"
    src = HERE / f"recheck_{day}.csv"
    if not src.exists():
        print(f"{src.name} 없음")
        return
    # (분 단위 시각, ad_type, keyword) → {device: (순위, 비고)}
    found = OrderedDict()
    for r in csv.DictReader(open(src, encoding="utf-8-sig")):
        if r["status"] != "노출" or r["ad_type"] not in SHEETS:
            continue
        note = f"재확인 · 첫 화면 {r['total_1page']}개 중"
        key = (r["checked_at"][:16], r["ad_type"], r["keyword"])
        found.setdefault(key, {})[r["device"]] = (int(r["rank"]), note)
    if not found:
        print("재확인 노출 전환 0건 — 엑셀 변경 없음")
        return

    from paths import XL_DIR
    path = XL_DIR / f"순위요약_{day[:4]}-{day[4:6]}.xlsx"
    wb = load_workbook(path)
    migrate(wb)                                   # 2026-10-13: '순위' 1시트 형식
    done = set()
    if SHEET in wb.sheetnames:
        for row in wb[SHEET].iter_rows(min_row=2, values_only=True):
            done.add((str(row[0]), TYPE_OF.get(row[1]), row[2]))
    new = [(k, v) for k, v in found.items() if k not in done]
    if not new:
        print("추가할 재확인 행 없음 (이미 반영됨)")
        return

    def write(book):
        for (ts, t, kw), devs in new:
            pc, mo = devs.get("pc"), devs.get("mobile")
            _append_styled(_sheet(book), [ts, SHEETS[t], kw, pc[0] if pc else "", mo[0] if mo else "",
                                          pc[1] if pc else "", mo[1] if mo else ""])

    write(wb)
    try:
        from summary import show_all
        show_all(wb)
        wb.save(path)
        print(f"{path.name} 에 재확인 {len(new)}행 추가")
        from publish_done import publish          # 2026-10-06: '모니터링완료_날짜_시각' 사본 갱신
        print("  " + publish(f"{day[:4]}-{day[4:6]}-{day[6:]}"))
    except PermissionError:
        alt = path.with_name(path.stem + f"_{datetime.now():%H%M}" + path.suffix)
        one = Workbook()
        one.remove(one.active)
        write(one)
        one.save(alt)
        print(f"[!] {path.name} 이 열려 있어 {alt.name} 으로 {len(new)}행 저장 (다음 정규 실행 때 자동 합침)")
    for (ts, t, kw), devs in new:
        print(" ", ts, SHEETS[t], kw, devs)


if __name__ == "__main__":
    main()
